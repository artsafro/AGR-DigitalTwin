from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/material-register-v001';OUT.mkdir(parents=True,exist_ok=True)
SOURCE=Path('C:/Users/artsafro/Downloads/Telegram Desktop/СОШ1150_СевЧертаново_АГР_24.09.2026 (1).pdf')
rows=[
 ('F1',['Ф.1','Ф.13'],'НВФ с облицовкой металлическими фасадными кассетами. Серый фасад и откосы F13.','NCS S 700-N¹',4096,'Текстура кассет','Код переписан буквально; требуется уточнение записи в альбоме.'),
 ('F2',['Ф.2'],'НВФ с облицовкой металлическими фасадными кассетами. Жёлтая отделка.','NCS S 2030-Y20R',256,'Заглушка','Однотонная карта.'),
 ('F3',['Ф.3'],'Декоративные фасадные элементы из металлических кассет на подсистеме. Фиолетовые.','NCS S 6020-R40B',256,'Заглушка','Однотонная карта.'),
 ('F4',['Ф.4'],'Декоративные фасадные элементы из металлических кассет на подсистеме. Светлые кремовые.','NCS S 0515-Y40R',256,'Заглушка','Однотонная карта.'),
 ('F5',['Ф.5'],'Декоративные фасадные элементы из металлических кассет на подсистеме. Синие.','NCS S 6020-R80B',256,'Заглушка','Однотонная карта.'),
 ('F6',['Ф.6'],'Декоративные фасадные элементы из металлических кассет на подсистеме. Красно-оранжевые.','NCS S 1060-Y90R',256,'Заглушка','Однотонная карта.'),
 ('F7',['Ф.7'],'Декоративные фасадные элементы из металлических кассет на подсистеме. Жёлтые.','NCS S 1070-Y10R',256,'Заглушка','Однотонная карта.'),
 ('F8',['Ф.8'],'Декоративные фасадные элементы из металлических кассет на подсистеме. Оранжевые.','NCS S 1060-Y40R',256,'Заглушка','Однотонная карта.'),
 ('F9',['Ф.9'],'Цоколь. Керамогранит «Уральский Гранит», моноколор УФ011, жёлтый.','УФ011\nЖёлтый',4096,'Плитка со швами','RAL в легенде не указан; цвет по артикулу производителя.'),
 ('F9.1',['Ф.9.1'],'Цоколь. Керамогранит «Уральский Гранит», моноколор UF002, светло-серый.','UF002\nСветло-серый',4096,'Плитка со швами','RAL в легенде не указан; цвет по артикулу производителя.'),
 ('F10',['Ф.10'],'Металлический декоративный элемент кровли, перголы.','NCS S 1060-Y40R\nNCS S 1060-Y90R²',256,'Заглушка','В альбоме указаны два цвета; единый цвет карты не выбран.'),
 ('F11',['Ф.11','Ф.14'],'Рамы и профили оконных блоков F11 и витражей F14. В альбоме: ПВХ и алюминий.','RAL 7016',256,'Заглушка','Объединение относится к рамам и профилям, не к прозрачному стеклу.'),
 ('F12',['Ф.12'],'Оконные отливы и фартук парапета над витражами. Оцинкованная сталь с полимерным покрытием.','RAL 7016',256,'Заглушка','По решению пользователя также прочий металл соответствующего цвета; F10/F16 не перекрашивать.'),
 ('F15',['Ф.15'],'Стемалитовые вставки в составе витражной системы.','RAL 7001',256,'Заглушка','Непрозрачное заполнение, отдельно от прозрачного стекла.'),
 ('F16',['Ф.16'],'Фартук парапета. Оцинкованная сталь с полимерным покрытием, заводское окрашивание.','RAL 1002',256,'Заглушка','Сохранять отдельно от тёмного металла F12.'),
]
data={'project':'СОШ на 1150 мест, Москва, мкр. Северное Чертаново, влд. 5',
      'source':str(SOURCE),'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
      'source_pdf_pages':[18,19,20,21],'date':'2026-09-29','aliases':{'F13':'F1','F14':'F11'},
      'resolution_source':'Прямое указание пользователя 29.09.2026','materials':[]}
for id,aliases,name,color,size,kind,note in rows:
    data['materials'].append({'id':id,'source_marks':aliases,'description':name,'source_color_code':color.replace('¹','').replace('²',''),
                              'resolution':[size,size],'resource_type':kind,'note':note,'color_source':'PDF legend pages 18-21',
                              'color_status':'NEEDS_CONFIRMATION' if id in ['F1','F10'] else 'DOCUMENTED'})
(OUT/'materials.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.7);sec.right_margin=Cm(1.7)
for sn in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[sn];st.font.name='Arial';st.font.color.rgb=RGBColor.from_string('000000')
doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].paragraph_format.space_after=Pt(9)
doc.styles['Subtitle'].font.size=Pt(11);doc.styles['Subtitle'].paragraph_format.space_after=Pt(11)
doc.styles['Subtitle'].font.italic=False
doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].paragraph_format.space_before=Pt(10);doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
doc.styles['Heading 2'].font.size=Pt(11.5);doc.styles['Heading 2'].font.bold=True
doc.core_properties.title='Реестр материалов фасадов СОШ на 1150 мест'
doc.core_properties.subject='Материалы F1–F16 по листам 18–21 и назначение размеров карт'
doc.core_properties.author='';doc.core_properties.keywords='Северное Чертаново, СОШ1150, материалы, NCS, RAL'

def p(text,style=None):return doc.add_paragraph(text,style)
def table(items):
    t=doc.add_table(rows=1, cols=4);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[1.05,8.0,4.3,4.25]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,w in zip(t.rows[0].cells,widths):c.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,['ID','Материал и назначение','Цвет по альбому','Карта в пикселях']):
        c.text=txt;shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECEF');c._tc.get_or_add_tcPr().append(shade)
        for r in c.paragraphs[0].runs:r.bold=True;r.font.size=Pt(10)
    trPr=t.rows[0]._tr.get_or_add_trPr();header=OxmlElement('w:tblHeader');trPr.append(header)
    for i,r in enumerate(items):
        id,aliases,name,color,size,kind,note=r;cells=t.add_row().cells
        for c,w in zip(cells,widths):c.width=Cm(w)
        values=[id,name,color,f'{size} × {size}\n{kind}']
        for j,(c,txt) in enumerate(zip(cells,values)):
            c.text=txt;c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1.04
                for run in para.runs:run.font.size=Pt(10);run.bold=j==0
            tcPr=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');marg.append(el)
            tcPr.append(marg)
            if i%2==1:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F5F6F7');tcPr.append(sh)
        pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D4DADF');borders.append(el)
    t._tbl.tblPr.append(borders)
    return t

p('Реестр материалов фасадов','Title')
p('СОШ на 1150 мест\nМосква, мкр. Северное Чертаново, влд. 5','Subtitle')
p('Для назначения отделок в модели и подготовки текстур. Приняты 15 рабочих ID: 3 текстуры 4096 × 4096 и 12 позиций с заглушками 256 × 256. F13 объединён с F1, F14 — с F11. Размеры карт заданы пользователем; названия и коды цветов сверены с листами 18–21 альбома АГР от 24.09.2026.')
p('Кассеты и декоративные фасадные элементы','Heading 1')
table(rows[:8])
p('¹ F1 и F13: в легенде буквально указано «NCS S 700-N». Запись нужно уточнить до подбора точного цвета. Исправленный код и приблизительный RAL не назначены.')
p('Для F2–F8 сохранены проектные обозначения NCS. Их не следует заменять приблизительными RAL по цвету визуализации.')

doc.add_page_break()
p('Цоколь оконные профили и металл','Heading 1')
table(rows[8:])
p('Правила назначения','Heading 1')
p('F13 → F1. Металлические откосы из фасадных кассет используют материал и текстуру F1, 4096 × 4096. Отдельную карту F13 не создавать.')
p('F14 → F11. Рамы и профили витражей используют F11, RAL 7016, 256 × 256. Прозрачное стекло остаётся отдельным заполнением; его ID в этом реестре не назначен. F15 сохраняется отдельно как стемалит.')
p('F9 и F9.1 — два отдельных ID и две текстуры 4096 × 4096 с плиткой и швами. Размер плитки, ширину и цвет швов зафиксировать перед созданием карт. В легенде указаны артикулы, а не RAL.')
p('² F10: в альбоме предусмотрены NCS S 1060-Y40R и NCS S 1060-Y90R. Одна однотонная заглушка не передаст оба цвета. Размер 256 × 256 сохраняется; распределение цветов по элементам требует уточнения. Эти коды уже встречаются у F8 и F6, но F10 автоматически к ним не объединён.')
p('F12 сохраняется отдельно от F11 при одинаковом RAL 7016: профиль окна и прочий металл различаются по назначению. Прочий металл относится к F12 только при том же цвете; F10 и F16 сохраняют собственные цвета.')
p('Источник','Heading 2')
pp=p('СОШ1150_СевЧертаново_АГР_24.09.2026 (1).pdf. Лист 18 — фасад 1–20; лист 19 — 20–1; лист 20 — А–НН; лист 21 — НН–А. Номера листов совпадают с номерами страниц PDF. Реестр составлен 29.09.2026. Заглушка означает однотонную карту; текстуры и шейдеры этим документом не создаются.')
for r in pp.runs:r.font.size=Pt(8.5)
for tree in [doc.styles.element,doc.element]:
    for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.save(OUT/'SOSH1150_Material_Register_v001.docx')
print(json.dumps({'file':str(OUT/'SOSH1150_Material_Register_v001.docx'),'ids':len(rows),'textures_4096':sum(r[4]==4096 for r in rows),'placeholders_256':sum(r[4]==256 for r in rows)},ensure_ascii=False))
